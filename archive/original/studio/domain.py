"""Composition validation and Excel import, independent of the desktop UI."""
from __future__ import annotations

from decimal import Decimal, InvalidOperation
from pathlib import Path
import re
from typing import Any


_ALIASES = {
    "Cotton": ["cotton", "棉", "棉花", "纯棉"],
    "Polyester": ["polyester", "polyester fiber", "polyester fibre", "聚酯纤维", "聚酯", "涤纶"],
    "Viscose": ["viscose", "rayon", "viscose rayon", "粘纤", "粘胶纤维", "黏胶纤维", "粘胶", "黏纤", "人造丝"],
    "Nylon": ["nylon", "polyamide", "锦纶", "尼龙"],
    "Elastane": ["elastane", "spandex", "lycra", "氨纶", "弹性纤维"],
    "Linen": ["linen", "flax", "亚麻"],
    "Ramie": ["ramie", "苎麻"],
    "Hemp": ["hemp", "大麻", "汉麻"],
    "Wool": ["wool", "羊毛"],
    "Cashmere": ["cashmere", "羊绒"],
    "Silk": ["silk", "桑蚕丝", "真丝", "蚕丝"],
    "Acrylic": ["acrylic", "腈纶"],
    "Modal": ["modal", "莫代尔", "莫代尔纤维"],
    "Lyocell": ["lyocell", "莱赛尔", "莱赛尔纤维", "莱塞尔", "天丝"],
    "Acetate": ["acetate", "醋纤", "醋酸纤维", "醋酯纤维"],
    "Triacetate": ["triacetate", "三醋酸纤维", "三醋酯纤维"],
    "Bamboo": ["bamboo", "bamboo fiber", "bamboo fibre", "竹纤维"],
    "Polypropylene": ["polypropylene", "丙纶"],
    "Cupro": ["cupro", "铜氨纤维", "铜氨丝"],
    "Mohair": ["mohair", "马海毛"],
    "Alpaca": ["alpaca", "羊驼毛"],
    "Metallic fiber": ["metallic fiber", "metallic fibre", "纯金属", "金属纤维"],
}
FIBERS = list(_ALIASES)
_NAME_MAP = {alias.casefold(): name for name, aliases in _ALIASES.items() for alias in aliases}
_NAME_MAP.update({name.casefold(): name for name in FIBERS})

COLORS = [
    {"id": "light-blue", "name": "Light blue / 浅蓝", "hex": "#B8C9D8"},
    {"id": "royal-blue", "name": "Royal blue / 宝蓝", "hex": "#315BBA"},
    {"id": "ivory", "name": "Ivory / 米白", "hex": "#DDDCC6"},
    {"id": "navy", "name": "Navy / 藏蓝", "hex": "#29283F"},
    {"id": "teal", "name": "Teal / 蓝绿", "hex": "#168294"},
    {"id": "black", "name": "Black / 黑色", "hex": "#232725"},
    {"id": "green", "name": "Green / 绿色", "hex": "#567846"},
    {"id": "burgundy", "name": "Burgundy / 酒红", "hex": "#572B40"},
    {"id": "white", "name": "White / 白色", "hex": "#F4F1E9"},
    {"id": "grey", "name": "Grey / 灰色", "hex": "#797979"},
    {"id": "red", "name": "Red / 红色", "hex": "#C20C25"},
    {"id": "orange", "name": "Orange / 橙色", "hex": "#F4A256"},
]

_NAMES = "|".join(re.escape(name) for name in sorted(_NAME_MAP, key=len, reverse=True))
_NUMBER = r"(?:\d+(?:\.\d+)?|\.\d+)"
_AFTER = re.compile(rf"(?P<name>{_NAMES})\s*[:：]?\s*(?P<percent>{_NUMBER})\s*%", re.IGNORECASE)
_BEFORE = re.compile(rf"(?P<percent>{_NUMBER})\s*%\s*(?P<name>{_NAMES})(?![A-Za-z])", re.IGNORECASE)
_SEPARATOR = re.compile(r"[\s,，;；/、+&·|]+")
_HEADERS = {"成分", "成份", "面料成分", "面料成份", "composition", "compositions", "fabric composition", "fiber content", "fibre content", "(空白)", "（空白）", "总计", "grand total"}


def _decimal(value: Any) -> Decimal:
    if isinstance(value, bool):
        raise ValueError("A percentage must be a number, not true/false.")
    try:
        number = Decimal(str(value).strip().rstrip("%"))
    except (InvalidOperation, ValueError, TypeError):
        raise ValueError(f"Invalid percentage: {value!r}.") from None
    if not number.is_finite() or number <= 0 or number > 100:
        raise ValueError("Each percentage must be greater than 0 and no more than 100.")
    return number


def validate_composition(parts: list[dict]) -> list[dict]:
    """Return normalized ordered parts, or raise ValueError; never guess fibres."""
    if not isinstance(parts, (list, tuple)) or not parts:
        raise ValueError("Add at least one fibre and percentage.")
    result: list[dict] = []
    names: set[str] = set()
    total = Decimal(0)
    for part in parts:
        if not isinstance(part, dict):
            raise ValueError("Each composition part must contain name and percent.")
        name = _NAME_MAP.get(str(part.get("name", "")).strip().casefold())
        if not name:
            raise ValueError(f"Unknown fibre: {part.get('name', '')!r}. Select a supported fibre explicitly.")
        if name in names:
            raise ValueError(f"{name} appears more than once; combine its percentages.")
        number = _decimal(part.get("percent"))
        names.add(name)
        total += number
        result.append({"name": name, "percent": float(number)})
    if abs(total - Decimal(100)) > Decimal("0.01"):
        raise ValueError(f"Percentages must total 100%; current total is {_format_percent(total)}%.")
    return result


def parse_composition(text: str) -> list[dict]:
    """Read e.g. 棉10%亚麻90% or '10% Cotton, 90% Linen', preserving order."""
    if not isinstance(text, str) or not text.strip():
        raise ValueError("Enter a composition, for example 棉100% or 100% Cotton.")
    text = text.replace("％", "%").replace("\u00a0", " ").strip()
    parts: list[dict] = []
    position = 0
    while position < len(text):
        separator = _SEPARATOR.match(text, position)
        if separator:
            position = separator.end()
            if position == len(text):
                break
        token = _AFTER.match(text, position) or _BEFORE.match(text, position)
        if not token:
            excerpt = text[position:position + 35]
            raise ValueError(f"Unrecognized composition text near {excerpt!r}. Use a fibre name with a percentage.")
        parts.append({"name": token.group("name"), "percent": token.group("percent")})
        position = token.end()
    return validate_composition(parts)


def _format_percent(number: Any) -> str:
    return format(Decimal(str(number)).normalize(), "f")


def composition_text(parts: list[dict]) -> str:
    """Exact label text; each part occupies one line, without a size line."""
    return "\n".join(f"{_format_percent(part['percent'])}% {part['name']}" for part in validate_composition(parts))


def import_column_c(path: str | Path, sheet: str | None = None) -> dict:
    """Read compositions from column C. Invalid rows remain visible as errors."""
    from openpyxl import load_workbook

    path = Path(path)
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("Choose an .xlsx or .xlsm workbook. Save older .xls files as .xlsx first.")
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheets = list(workbook.sheetnames)
        if sheet is not None and sheet not in sheets:
            raise ValueError(f"Worksheet {sheet!r} was not found.")
        worksheet = workbook[sheet] if sheet else workbook.active
        compositions, errors = [], []
        for row in worksheet.iter_rows(min_col=3, max_col=3):
            cell = row[0]
            value = cell.value
            if value is None or not str(value).strip():
                continue
            raw = str(value).strip()
            if raw.casefold() in _HEADERS:
                continue
            try:
                if cell.data_type == "f":
                    raise ValueError("Column C contains a formula. Paste its composition result as text before importing.")
                parts = parse_composition(raw)
                compositions.append({"source": f"{worksheet.title}!C{cell.row}", "label": composition_text(parts).replace("\n", " · "), "parts": parts})
            except ValueError as error:
                errors.append({"row": cell.row, "value": raw, "error": str(error)})
        return {"compositions": compositions, "errors": errors, "sheets": sheets}
    finally:
        workbook.close()
